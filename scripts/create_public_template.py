"""Create a portable contest template containing fictitious examples only."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def build_template():
    book = Workbook()
    result = book.active
    result.title = 'Result'
    criteria = book.create_sheet('Criteria')
    help_sheet = book.create_sheet('Petunjuk')
    values = {
        'A2': 'TARGET STORES', 'J2': 'TARGET AM', 'O2': 'TARGET RSM/Asst. RSM',
        'A3': 'No', 'B3': 'Store ID', 'C3': 'Store Name', 'D3': 'Store Size', 'E3': 'Target Revenue',
        'F3': 'Reward Stores', 'H3': 'Scheme', 'J3': 'AM', 'K3': 'Store Area', 'L3': 'Target 90%',
        'M3': 'Target 100%', 'O3': 'RSM/Asst. RSM', 'P3': 'Target 90%', 'Q3': 'Target 100%',
        'H4': 'CONTOH FIKTIF. Ganti skema kategori Anda. Target produk: LAPTOP (ALL BRAND). Status: New & Display (Sample). Reward toko: minimal 100% target, ranking menurut SO/target. Periksa filter dan aturan di web.',
        'F7': 'Rank 1', 'G7': 1000000, 'F8': 'Rank 2-10', 'G8': 500000,
        'J4': 'AM Contoh', 'K4': 2, 'L4': 18000000, 'M4': 20000000,
        'O4': 'RSM Contoh', 'P4': 18000000, 'Q4': 20000000,
        'O10': 'Reward RSM', 'O12': 'RSM Contoh', 'O14': 'Achieve 100%', 'Q14': 1000000,
        'O15': 'Achieve 90%', 'Q15': 500000, 'J27': 'Reward AM', 'J29': 'Target Achieve',
        'K29': 'Target 90%', 'L29': 'Target 100%',
        'J30': 'Jika Jumlah Toko yang Ditangani AM ≤ 7 Toko', 'K30': 250000, 'L30': 500000,
        'J31': 'Jika Jumlah Toko yang Ditangani AM ≥ 8 Toko', 'K31': 500000, 'L31': 1000000,
    }
    for cell, value in values.items(): criteria[cell] = value
    result['A2'] = 'RESULT SALES CONTEST NAMA KATEGORI'
    for cell, value in {'A3': 'NO', 'B3': 'AM', 'C3': 'ID Store', 'D3': 'Store Name', 'E3': 'Target',
                        'F3': 'SO', 'G3': '%', 'H3': 'Reward', 'J3': 'AM', 'K3': 'Store Area',
                        'L3': 'Target 90%', 'M3': 'Target AM 100%', 'N3': 'SO', 'O3': '% Target 90%',
                        'P3': '% Target 100%', 'Q3': 'Rank', 'R3': 'Rewards', 'T3': 'RSM',
                        'U3': 'Target 90%', 'V3': 'Target 100%', 'W3': 'SO', 'X3': '% Target 90%',
                        'Y3': '% Target 100%', 'Z3': 'Reward RSM'}.items(): result[cell] = value
    for row in (4, 5):
        code, name = f'CONTOH00{row-3}', f'Toko Contoh {row-3}'
        for col, value in enumerate([row-3, code, name, 'M', 10000000], 1): criteria.cell(row, col, value)
        for col, value in enumerate([row-3, 'AM Contoh', code, name, 10000000], 1): result.cell(row, col, value)
    for text in [
        'PETUNJUK: seluruh nama dan angka adalah contoh fiktif, bukan data contest aktual.',
        'Satu file untuk satu kategori. Ganti judul Result dan nama kategori pada web.',
        'Ganti/hapus baris CONTOH. Isi Criteria B:E mulai baris 4 untuk toko peserta. Target harus numerik.',
        'Isi Criteria J:M untuk target AM, O:Q baris 4-7 untuk target RSM. Ketentuan reward di bawah tabel.',
        'Isi identitas Result B:D sesuai toko, atau upload mapping AM/RSM terpisah.',
        'SO dan bonus dihitung aplikasi dari Revenue. Periksa filter dan aturan di web sebelum menghitung.',
        'Pencocokan berdasarkan Store ID. Nama AM/RSM harus sesuai mapping.',
    ]: help_sheet.append([text])
    for sheet in (result, criteria):
        sheet.freeze_panes = 'E4'
        for cell in sheet[3]:
            if cell.value:
                cell.fill = PatternFill('solid', fgColor='0066B3')
                cell.font = Font(bold=True, color='FFFFFF')
                cell.alignment = Alignment(wrap_text=True)
        for col in range(1, sheet.max_column + 1): sheet.column_dimensions[sheet.cell(1, col).column_letter].width = 21
        sheet.row_dimensions[3].height = 32
    criteria.column_dimensions['C'].width = 32
    criteria.column_dimensions['H'].width = 60
    criteria.column_dimensions['J'].width = 45
    criteria['H4'].alignment = Alignment(wrap_text=True)
    criteria.row_dimensions[4].height = 95
    help_sheet.column_dimensions['A'].width = 115
    for row in help_sheet:
        row[0].alignment = Alignment(wrap_text=True)
        help_sheet.row_dimensions[row[0].row].height = 32
    target = Path(__file__).resolve().parents[1] / 'assets' / 'Template Kriteria Sales Contest.xlsx'
    book.save(target)
    return target


if __name__ == '__main__':
    print(build_template())
