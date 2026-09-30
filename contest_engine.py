"""Excel input, configurable contest calculations, and value-based exports."""
from __future__ import annotations

from copy import copy
from io import BytesIO
import re

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment


def clean(value):
    if value is None or pd.isna(value):
        return ''
    return re.sub(r'\s+', ' ', str(value)).strip()


def identity(value):
    return re.sub(r'[^\w]', '', clean(value).casefold())


def store_id(value):
    value = clean(value)
    return value[:-2] if re.fullmatch(r'\d+\.0', value) else value


def number(value):
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def read_template(content):
    book = openpyxl.load_workbook(BytesIO(content), data_only=True)
    names = {s.casefold().strip(): s for s in book.sheetnames}
    if not {'criteria', 'result'} <= names.keys():
        raise ValueError('Template harus memiliki sheet Criteria dan Result.')
    criteria, result = book[names['criteria']], book[names['result']]
    # Store/manager blocks follow the supplied template, joined by ID rather than row.
    if 'store' not in clean(criteria['B3'].value).lower():
        raise ValueError('Format Criteria tidak sesuai template: Store ID harus berada di B3.')
    stores = []
    for row in criteria.iter_rows(min_row=4, max_col=17, values_only=True):
        if clean(row[1]) and clean(row[2]) and number(row[4]) is None:
            raise ValueError(f'Target toko {store_id(row[1])} kosong/tidak numerik. Jika memakai formula, simpan ulang file melalui Excel agar hasil formula tersedia.')
        if clean(row[1]) and clean(row[2]) and number(row[4]) is not None:
            stores.append({'Store ID': store_id(row[1]), 'Store Name': clean(row[2]),
                           'Target': number(row[4])})
    if not stores:
        raise ValueError('Criteria belum memiliki target toko numerik.')
    stores = pd.DataFrame(stores)
    if stores['Store ID'].duplicated().any():
        raise ValueError('Store ID duplikat pada Criteria. Satu toko hanya boleh memiliki satu target.')
    if (stores.Target < 0).any():
        raise ValueError('Target toko tidak boleh negatif.')
    mapping = []
    for row in result.iter_rows(min_row=4, max_col=4, values_only=True):
        if clean(row[2]) and clean(row[1]) and store_id(row[2]) in set(stores['Store ID']):
            mapping.append({'Store ID': store_id(row[2]), 'Store Name': clean(row[3]),
                            'AM': clean(row[1]), 'RSM': ''})
    am, rsm = [], []
    for row in criteria.iter_rows(min_row=4, max_col=17, values_only=True):
        if clean(row[9]) and number(row[12]) is not None and number(row[11]) is not None:
            am.append({'AM': clean(row[9]), 'Jumlah Toko': number(row[10]) or 0,
                       'Target 90%': number(row[11]), 'Target 100%': number(row[12])})
        if clean(row[14]) and number(row[15]) is not None and number(row[16]) is not None:
            rsm.append({'RSM': clean(row[14]), 'Target 90%': number(row[15]),
                        'Target 100%': number(row[16]), 'Reward 90%': 0., 'Reward 100%': 0.})
    ranks = []
    for row in criteria.iter_rows(min_row=4, max_col=7, values_only=True):
        match = re.search(r'rank\s*(\d+)(?:\s*[-–]\s*(\d+))?', clean(row[5]), re.I)
        if match and number(row[6]) is not None:
            ranks.append({'Dari Rank': int(match[1]), 'Sampai Rank': int(match[2] or match[1]),
                          'Reward': number(row[6])})
    am_rules = []
    for row in criteria.iter_rows(min_row=4, max_col=12, values_only=True):
        label = clean(row[9])
        if 'toko' in label.casefold() and number(row[10]) is not None and number(row[11]) is not None:
            match = re.search(r'([≤≥<>]=?)\s*(\d+)', label)
            if match:
                op, limit = match[1], int(match[2])
                low, high = (0, limit) if op in ('≤', '<=') else (limit, 100000)
                if op == '<': low, high = 0, limit - 1
                if op == '>': low, high = limit + 1, 100000
                am_rules.append({'Min Toko': low, 'Max Toko': high,
                                 'Reward 90%': number(row[10]), 'Reward 100%': number(row[11])})
    recipients = []
    for row in criteria.iter_rows(min_row=10, max_col=17, values_only=True):
        label = clean(row[14])
        if not label:
            continue
        matching = [r for r in rsm if identity(r['RSM']) in identity(label)]
        if matching:
            recipients = matching
        if 'achieve' in label.casefold() and number(row[16]) is not None:
            key = 'Reward 100%' if '100%' in label else 'Reward 90%' if '90%' in label else None
            if key:
                for r in recipients:
                    r[key] = number(row[16])
    title = clean(result['A2'].value)
    category = re.sub(r'^RESULT SALES CONTEST\s*', '', title, flags=re.I)
    category = re.sub(r'\s*-\s*\w+\s+20\d\d.*$', '', category).strip() or 'Kategori'
    output = dict(stores=stores, mapping=pd.DataFrame(mapping, columns=['Store ID', 'Store Name', 'AM', 'RSM']),
                  am=pd.DataFrame(am, columns=['AM', 'Jumlah Toko', 'Target 90%', 'Target 100%']),
                  rsm=pd.DataFrame(rsm, columns=['RSM', 'Target 90%', 'Target 100%', 'Reward 90%', 'Reward 100%']),
                  ranks=pd.DataFrame(ranks, columns=['Dari Rank', 'Sampai Rank', 'Reward']),
                  am_rules=pd.DataFrame(am_rules, columns=['Min Toko', 'Max Toko', 'Reward 90%', 'Reward 100%']),
                  scheme=clean(criteria['H4'].value), category=category)
    book.close()
    return output


REVENUE_COLUMNS = ['Date', 'Store code', 'Store name', 'Main group', 'Sub group',
                   'Product code', 'Product name', 'Brand name', 'Inventory status',
                   'Out put type name', 'Quantity', 'Revenue']


def read_revenue(content):
    book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = book.active
    header = [clean(v) for v in next(sheet.iter_rows(max_row=1, values_only=True))]
    missing = set(REVENUE_COLUMNS) - set(header)
    if missing:
        raise ValueError('Kolom Revenue tidak ditemukan: ' + ', '.join(sorted(missing)))
    positions = [header.index(c) for c in REVENUE_COLUMNS]
    values = []
    for row in sheet.iter_rows(min_row=2, max_col=max(positions) + 1, values_only=True):
        if any(v is not None for v in row):
            values.append([row[i] for i in positions])
    book.close()
    frame = pd.DataFrame(values, columns=REVENUE_COLUMNS)
    frame['Store code'] = frame['Store code'].map(store_id)
    for col in REVENUE_COLUMNS[2:10]:
        frame[col] = frame[col].map(clean)
    for col in ('Revenue', 'Quantity'):
        parsed = pd.to_numeric(frame[col], errors='coerce')
        if parsed.isna().any():
            raise ValueError(f'{col}: {int(parsed.isna().sum())} baris kosong/tidak numerik. Perbaiki file terlebih dahulu.')
        frame[col] = parsed
    frame['Date'] = pd.to_datetime(frame['Date'], dayfirst=True, format='mixed', errors='coerce')
    if frame['Date'].isna().any() or frame['Store code'].eq('').any():
        raise ValueError('Revenue memiliki tanggal atau kode toko kosong/tidak valid.')
    return frame


def validate_mapping(frame):
    frame = frame.copy()
    for col in ['Store ID', 'AM', 'RSM']:
        frame[col] = frame[col].map(store_id if col == 'Store ID' else clean)
    frame = frame[frame['Store ID'] != ''].drop_duplicates(['Store ID', 'AM', 'RSM'])
    if frame.empty:
        raise ValueError('Mapping toko belum diisi.')
    # A store may belong to multiple RSM/assistant recipients, but only one AM.
    counts = frame.assign(_am=frame.AM.map(identity)).groupby('Store ID')._am.nunique()
    if (counts > 1).any():
        raise ValueError('Satu toko dipetakan ke beberapa AM: ' + ', '.join(counts[counts > 1].index[:8]))
    return frame


FILTER_COLUMNS = ['Main group', 'Sub group', 'Brand name', 'Inventory status', 'Out put type name']


def filter_revenue(frame, config):
    selected = frame[(frame.Date.dt.date >= config['start']) & (frame.Date.dt.date <= config['end'])]
    for col, values in config['filters'].items():
        if values:
            selected = selected[selected[col].isin(values)]
    return selected


def validate_rules(config):
    for table, lower, upper in [(config['ranks'], 'Dari Rank', 'Sampai Rank'),
                                 (config['am_rules'], 'Min Toko', 'Max Toko')]:
        if table.empty:
            continue
        if table.isna().any().any():
            raise ValueError('Tabel aturan memiliki sel kosong.')
        for col in table.columns:
            if (pd.to_numeric(table[col], errors='coerce').isna() | (table[col] < 0)).any():
                raise ValueError('Nilai aturan harus berupa angka nonnegatif.')
        if ((table[lower] % 1 != 0) | (table[upper] % 1 != 0)).any():
            raise ValueError('Batas ranking/jumlah toko harus bilangan bulat.')
        if lower == 'Dari Rank' and (table[lower] < 1).any():
            raise ValueError('Ranking dimulai dari 1.')
        previous = -1
        for _, row in table.sort_values(lower).iterrows():
            if row[lower] > row[upper] or row[lower] <= previous:
                raise ValueError('Rentang aturan tumpang tindih atau batas awal melebihi batas akhir.')
            previous = row[upper]
    for name in ['am', 'rsm']:
        table = config[name]
        who = name.upper()
        if not table.empty:
            if table[who].map(identity).eq('').any() or table[who].map(identity).duplicated().any():
                raise ValueError(f'Nama {who} kosong atau duplikat.')
            for col in table.columns.drop(who):
                vals = pd.to_numeric(table[col], errors='coerce')
                if vals.isna().any() or (vals < 0).any():
                    raise ValueError(f'{who}: {col} harus angka nonnegatif.')
            if (table['Target 90%'] > table['Target 100%']).any():
                raise ValueError(f'{who}: target 90% tidak boleh lebih besar dari target 100%.')


def calculate(template, revenue, mapping, config):
    validate_rules(config)
    mapping = validate_mapping(mapping)
    selected = filter_revenue(revenue, config)
    so = selected.groupby('Store code')[config['metric']].sum()
    stores = template['stores'].copy()
    one_am = mapping.drop_duplicates('Store ID').set_index('Store ID').AM
    stores['AM'] = stores['Store ID'].map(one_am).fillna('')
    if stores.AM.eq('').any():
        missing = stores.loc[stores.AM.eq(''), 'Store ID'].tolist()
        raise ValueError('Mapping AM belum lengkap untuk toko peserta: ' + ', '.join(missing[:10]))
    stores['SO'] = stores['Store ID'].map(so).fillna(0.)
    stores['Pencapaian'] = stores['SO'].div(stores.Target.where(stores.Target > 0))
    stores['Status'] = stores.Target.map(lambda t: 'OK' if t > 0 else 'Target nol — tidak ikut ranking')
    stores = stores.sort_values(['Pencapaian', 'Store ID'], ascending=[False, True], kind='stable', na_position='last')
    eligible = (stores.Target > 0) & (stores.Pencapaian >= config['minimum'] / 100)
    stores['Rank'] = pd.Series(pd.NA, index=stores.index, dtype='Int64')
    if config['ties'] == 'Urutan Store ID':
        stores.loc[eligible, 'Rank'] = range(1, int(eligible.sum()) + 1)
    else:
        stores.loc[eligible, 'Rank'] = stores.loc[eligible, 'Pencapaian'].rank(method='min', ascending=False).astype('Int64')
    stores['Reward'] = 0. if not config['ranks'].empty else float('nan')
    for _, rule in config['ranks'].iterrows():
        hit = stores.Rank.between(rule['Dari Rank'], rule['Sampai Rank']).fillna(False)
        stores.loc[hit, 'Reward'] = rule['Reward']
    notes = []
    am_rows = []
    for _, target in config['am'].iterrows():
        manager = target.AM
        member_ids = set(mapping.loc[mapping.AM.map(identity) == identity(manager), 'Store ID'])
        if config['manager_scope'] == 'Toko peserta kategori':
            member_ids &= set(stores['Store ID'])
        amount = float(so.reindex(list(member_ids)).fillna(0).sum())
        count = len(member_ids) if config['count_source'] == 'Mapping' else int(target['Jumlah Toko'])
        reward = 0.
        covered = any(rule['Min Toko'] <= count <= rule['Max Toko'] for _, rule in config['am_rules'].iterrows())
        valid = bool(member_ids) and target['Target 100%'] > 0 and covered
        for _, rule in config['am_rules'].iterrows():
            if valid and rule['Min Toko'] <= count <= rule['Max Toko']:
                if amount >= target['Target 100%']:
                    reward = rule['Reward 100%']
                elif target['Target 90%'] > 0 and amount >= target['Target 90%']:
                    reward = rule['Reward 90%']
        am_rows.append({'AM': manager, 'Jumlah Toko': count, 'Target': target['Target 100%'],
                        'SO': amount, 'Pencapaian': amount / target['Target 100%'] if target['Target 100%'] > 0 else None,
                        'Reward': reward if valid else None, 'Status': 'OK' if valid else 'Mapping/target belum lengkap'})
    am = pd.DataFrame(am_rows, columns=['AM', 'Jumlah Toko', 'Target', 'SO', 'Pencapaian', 'Reward', 'Status'])
    missing_am = set(stores.AM.map(identity)) - set(am.AM.map(identity))
    if missing_am:
        notes.append(f'{len(missing_am)} AM toko peserta belum memiliki target/aturan AM; bonus AM belum lengkap.')
        for name in stores.AM.unique():
            if identity(name) in missing_am:
                am.loc[len(am)] = [name, None, None, None, None, None, 'Target AM belum tersedia']
    rsm_rows = []
    for _, target in config['rsm'].iterrows():
        member_ids = set(mapping.loc[mapping.RSM.map(identity) == identity(target.RSM), 'Store ID'])
        if config['manager_scope'] == 'Toko peserta kategori':
            member_ids &= set(stores['Store ID'])
        valid = bool(member_ids) and target['Target 100%'] > 0
        amount = float(so.reindex(list(member_ids)).fillna(0).sum()) if member_ids else None
        reward = None if not valid else (target['Reward 100%'] if amount >= target['Target 100%'] else
                  target['Reward 90%'] if target['Target 90%'] > 0 and amount >= target['Target 90%'] else 0.)
        rsm_rows.append({'RSM': target.RSM, 'Jumlah Toko': len(member_ids), 'Target': target['Target 100%'],
                         'SO': amount, 'Pencapaian': amount / target['Target 100%'] if valid else None,
                         'Reward': reward, 'Status': 'OK' if valid else 'Mapping/target belum lengkap'})
    rsm = pd.DataFrame(rsm_rows, columns=['RSM', 'Jumlah Toko', 'Target', 'SO', 'Pencapaian', 'Reward', 'Status'])
    missing_rsm = set(mapping.loc[mapping.RSM != '', 'RSM'].map(identity)) - set(rsm.RSM.map(identity))
    for name in mapping.RSM.unique():
        if identity(name) in missing_rsm:
            rsm.loc[len(rsm)] = [name, None, None, None, None, None, 'Target RSM belum tersedia']
    rsm_missing = set(stores['Store ID']) - set(mapping.loc[mapping.RSM != '', 'Store ID'])
    if rsm_missing:
        notes.append(f'{len(rsm_missing)} toko peserta belum memiliki mapping RSM; hasil RSM belum lengkap.')
    if not config['ranks'].shape[0]: notes.append('Aturan reward toko belum diisi.')
    if not config['am_rules'].shape[0]: notes.append('Aturan reward AM belum diisi.')
    if config['rsm'].empty: notes.append('Target dan reward RSM belum diisi.')
    if (stores.Target <= 0).any(): notes.append('Terdapat target toko nol; toko tersebut tidak ikut ranking.')
    if selected.empty: notes.append('Tidak ada transaksi yang memenuhi filter dan periode kategori ini.')
    for label, table in [('AM', am), ('RSM', rsm)]:
        if not table.empty and (table.Status != 'OK').any():
            notes.append(f'{label}: terdapat penerima tanpa mapping/target lengkap; bonus ditampilkan kosong.')
    without_sales = set(stores['Store ID']) - set(revenue.loc[(revenue.Date.dt.date >= config['start']) & (revenue.Date.dt.date <= config['end']), 'Store code'])
    if without_sales:
        notes.append(f'{len(without_sales)} toko peserta tidak ditemukan dalam Revenue periode ini; SO = 0.')
    return {'name': config['name'], 'stores': stores.reset_index(drop=True), 'am': am, 'rsm': rsm,
            'notes': notes, 'rows': len(selected), 'config': config}


def resume(results, level):
    key = {'stores': 'Store ID', 'am': 'AM', 'rsm': 'RSM'}[level]
    frames = []
    for result in results:
        frame = result[level][[key, 'Reward']].copy()
        frame['_key'] = frame[key].map(store_id if level == 'stores' else identity)
        frame = frame.rename(columns={'Reward': result['name']}).set_index('_key')
        frames.append(frame)
    out = pd.concat([frame.drop(columns=key) for frame in frames], axis=1)
    labels = pd.concat([frame[key] for frame in frames])
    labels = labels[~labels.index.duplicated()]
    out.insert(0, key, labels.reindex(out.index))
    if out.empty: return pd.DataFrame(columns=[key, 'Total'])
    categories = [r['name'] for r in results]
    # Absence from a category is zero; explicitly uncomputed rewards stay blank.
    for result in results:
        present = set(result[level][key].map(store_id if level == 'stores' else identity))
        out.loc[~out.index.isin(present), result['name']] = 0.
    out['Total'] = out[categories].sum(axis=1, min_count=len(categories))
    if level == 'stores':
        details = pd.concat([r['stores'][['Store ID', 'Store Name', 'AM']] for r in results]).drop_duplicates('Store ID').set_index('Store ID')
        out.insert(1, 'Store Name', out[key].map(details['Store Name']))
        out.insert(2, 'AM', out[key].map(details.AM))
    return out.reset_index(drop=True)


def export_results(results, templates, period):
    book = openpyxl.Workbook()
    book.remove(book.active)
    used = set()

    def new_sheet(name):
        name = re.sub(r'[\\/*?:\[\]]', '-', name)[:31] or 'Sheet'
        stem, n = name, 2
        while name.casefold() in used:
            name = stem[:26] + f' ({n})'
            n += 1
        used.add(name.casefold())
        return book.create_sheet(name)

    def write_frame(sheet, frame, row=1, col=1):
        for j, name in enumerate(frame.columns, col):
            cell = sheet.cell(row, j, name)
            cell.fill = PatternFill('solid', fgColor='123C40')
            cell.font = Font(color='FFFFFF', bold=True)
        for i, values in enumerate(frame.itertuples(index=False, name=None), row + 1):
            for j, value in enumerate(values, col):
                if pd.isna(value): value = None
                if hasattr(value, 'item'): value = value.item()
                cell = sheet.cell(i, j, value)
                if isinstance(value, str): cell.data_type = 's'
                if frame.columns[j-col] == 'Pencapaian': cell.number_format = '0.00%'
                elif isinstance(value, (int, float)): cell.number_format = '#,##0.00'
        for j, name in enumerate(frame.columns, col):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 38 if name in ('Store Name', 'AM', 'RSM', 'Status') else 22
        sheet.freeze_panes = 'A2'

    for level, title in [('stores', 'Resume Toko'), ('am', 'Resume AM'), ('rsm', 'Resume RSM')]:
        sheet = new_sheet(title)
        write_frame(sheet, resume(results, level))
        sheet.auto_filter.ref = sheet.dimensions
    notes = new_sheet('Catatan')
    notes.append(['Periode', period])
    notes.append(['Keterangan', 'Nilai dihitung aplikasi; sel bonus kosong berarti data belum lengkap.'])
    notes.append(['Kategori', 'Catatan / parameter'])
    for result in results:
        config = result['config']
        notes.append([result['name'], f"Metrik {config['metric']}; minimum toko {config['minimum']}%; ranking {config['ties']}; cakupan {config['manager_scope']}; jumlah toko AM: {config['count_source']}"])
        for field, values in config['filters'].items():
            notes.append([result['name'], f'{field}: ' + (', '.join(values) if values else 'Semua')])
        for note in result['notes']: notes.append([result['name'], note])
        sheet = new_sheet('Result ' + result['name'])
        write_frame(sheet, result['stores'])
        write_frame(sheet, result['am'], col=12)
        write_frame(sheet, result['rsm'], col=21)
        rules = new_sheet('Aturan ' + result['name'])
        offset = 1
        for key in ['ranks', 'am_rules', 'am', 'rsm']:
            rules.cell(offset, 1, key)
            write_frame(rules, config[key], row=offset + 1)
            offset += len(config[key]) + 4
    # Preserve uploaded Criteria for traceability; calculated Result never uses stale caches.
    for result, content in zip(results, templates):
        source = openpyxl.load_workbook(BytesIO(content), data_only=False)
        criteria = next(s for s in source if s.title.strip().casefold() == 'criteria')
        dest = new_sheet('Criteria ' + result['name'])
        for row in criteria:
            for cell in row:
                target = dest.cell(cell.row, cell.column, cell.value)
                if cell.has_style:
                    target.font = copy(cell.font)
                    target.fill = copy(cell.fill)
                    target.border = copy(cell.border)
                    target.alignment = copy(cell.alignment)
                    target.protection = copy(cell.protection)
                target.number_format = cell.number_format
        for merged in criteria.merged_cells.ranges: dest.merge_cells(str(merged))
        for key, dimension in criteria.column_dimensions.items(): dest.column_dimensions[key].width = dimension.width
        source.close()
    for row in notes:
        for cell in row:
            if isinstance(cell.value, str): cell.data_type = 's'
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    notes.column_dimensions['A'].width = 28
    notes.column_dimensions['B'].width = 110
    output = BytesIO()
    book.save(output)
    return output.getvalue()
