import unittest
from copy import deepcopy
from datetime import date
from pathlib import Path
from io import BytesIO

import openpyxl
import pandas as pd
from contest_engine import calculate, resume, export_results, read_template, validate_mapping

ROOT = Path(__file__).resolve().parents[1]


class CalculationTests(unittest.TestCase):
    def setUp(self):
        self.template = {'stores': pd.DataFrame([
            {'Store ID': '1', 'Store Name': 'One', 'Target': 100.},
            {'Store ID': '2', 'Store Name': 'Two', 'Target': 100.},
            {'Store ID': '3', 'Store Name': 'Zero target', 'Target': 0.}])}
        self.mapping = pd.DataFrame([
            {'Store ID': '1', 'AM': 'AM A', 'RSM': 'RSM A'},
            {'Store ID': '1', 'AM': 'AM A', 'RSM': 'RSM B'},
            {'Store ID': '2', 'AM': 'AM A', 'RSM': 'RSM A'},
            {'Store ID': '3', 'AM': 'AM A', 'RSM': 'RSM A'},
            {'Store ID': '4', 'AM': 'AM A', 'RSM': 'RSM A'}])
        self.revenue = pd.DataFrame([
            {'Date': pd.Timestamp('2026-08-01'), 'Store code': '1', 'Revenue': 120., 'Brand name': 'A'},
            {'Date': pd.Timestamp('2026-08-01'), 'Store code': '2', 'Revenue': 120., 'Brand name': 'A'},
            {'Date': pd.Timestamp('2026-08-01'), 'Store code': '4', 'Revenue': 50., 'Brand name': 'A'},
            {'Date': pd.Timestamp('2026-09-01'), 'Store code': '1', 'Revenue': 9999., 'Brand name': 'A'},
            {'Date': pd.Timestamp('2026-08-01'), 'Store code': '1', 'Revenue': 9999., 'Brand name': 'B'}])
        self.config = {'name': 'Kategori bebas', 'start': date(2026, 8, 1), 'end': date(2026, 8, 31),
            'filters': {'Brand name': ['A']}, 'metric': 'Revenue', 'minimum': 100.,
            'ties': 'Peringkat sama (1, 1, 3)', 'manager_scope': 'Toko peserta kategori', 'count_source': 'Mapping',
            'ranks': pd.DataFrame([{'Dari Rank': 1, 'Sampai Rank': 1, 'Reward': 20.}, {'Dari Rank': 2, 'Sampai Rank': 2, 'Reward': 10.}]),
            'am_rules': pd.DataFrame([{'Min Toko': 0, 'Max Toko': 7, 'Reward 90%': 5., 'Reward 100%': 10.}]),
            'am': pd.DataFrame([{'AM': 'AM A', 'Jumlah Toko': 3, 'Target 90%': 225., 'Target 100%': 250.}]),
            'rsm': pd.DataFrame([{'RSM': 'RSM A', 'Target 90%': 225., 'Target 100%': 250., 'Reward 90%': 15., 'Reward 100%': 30.},
                                 {'RSM': 'RSM B', 'Target 90%': 90., 'Target 100%': 100., 'Reward 90%': 10., 'Reward 100%': 20.}])}

    def run_calc(self):
        return calculate(self.template, self.revenue, self.mapping, self.config)

    def test_filters_ties_and_no_double_counting(self):
        r = self.run_calc()
        self.assertEqual(r['stores'].SO.sum(), 240.)
        self.assertEqual(r['stores'].Reward.sum(), 40.)
        self.assertEqual(r['am'].iloc[0].SO, 240.)
        self.assertEqual(r['am'].iloc[0].Reward, 5.)
        self.assertEqual(r['rsm'].Reward.sum(), 35.)
        self.assertTrue(pd.isna(r['stores'].iloc[2].Rank))

    def test_unique_rank_and_scope(self):
        self.config['ties'] = 'Urutan Store ID'
        self.config['manager_scope'] = 'Seluruh toko dalam mapping'
        r = self.run_calc()
        self.assertEqual(r['stores'].Reward.sum(), 30.)
        self.assertEqual(r['am'].iloc[0].SO, 290.)
        self.assertEqual(r['am'].iloc[0].Reward, 10.)

    def test_dynamic_categories_resume_am_once(self):
        first = self.run_calc()
        self.config['name'] = 'Kategori kedua'
        second = self.run_calc()
        summary = resume([first, second], 'am')
        self.assertEqual(len(summary), 1)
        self.assertEqual(summary.iloc[0].Total, 10.)

    def test_missing_mapping_is_blank_not_zero(self):
        self.mapping['RSM'] = ''
        r = self.run_calc()
        self.assertTrue(r['rsm'].Reward.isna().all())
        self.assertTrue(resume([r], 'rsm').Total.isna().all())

    def test_invalid_mapping_and_overlapping_rules(self):
        bad = pd.concat([self.mapping, pd.DataFrame([{'Store ID': '1', 'AM': 'Other', 'RSM': 'RSM A'}])])
        with self.assertRaises(ValueError): validate_mapping(bad)
        self.config['ranks'].loc[1, 'Dari Rank'] = 1
        with self.assertRaises(ValueError): self.run_calc()

    @unittest.skipUnless((ROOT / 'Template Kriteria Upload Sales Contes.xlsx').exists(), 'Original local sample is not distributed')
    def test_export_roundtrip_and_original_template(self):
        content = (ROOT / 'Template Kriteria Upload Sales Contes.xlsx').read_bytes()
        t = read_template(content)
        self.assertEqual(len(t['stores']), 261)
        self.assertEqual(len(t['am']), 18)
        self.assertEqual(len(t['rsm']), 4)
        self.assertEqual(t['rsm']['Reward 100%'].sum(), 5500000.)
        r = self.run_calc()
        exported = export_results([r], [content], 'Agustus 2026')
        book = openpyxl.load_workbook(BytesIO(exported), data_only=True)
        self.assertIn('Resume AM', book.sheetnames)
        self.assertEqual(book['Resume AM']['C2'].value, 5.)
        source = openpyxl.load_workbook(BytesIO(content), data_only=True)
        self.assertEqual(book['Criteria Kategori bebas']['H4'].value, source['Criteria']['H4'].value)
        self.assertTrue(book['Criteria Kategori bebas']['A3'].font.bold)

    def test_absent_category_different_recipients(self):
        first = self.run_calc()
        second = deepcopy(first)
        second['name'] = 'Another'
        second['am']['AM'] = 'Other manager'
        summary = resume([first, second], 'am')
        self.assertEqual(len(summary), 2)
        self.assertTrue((summary.Total == 5.).all())

    def test_empty_rules_leave_rewards_blank(self):
        self.config['ranks'] = self.config['ranks'].iloc[:0]
        self.config['am_rules'] = self.config['am_rules'].iloc[:0]
        r = self.run_calc()
        self.assertTrue(r['stores'].Reward.isna().all())
        self.assertTrue(r['am'].Reward.isna().all())

    def test_resume_first_category_has_no_rsm(self):
        first = self.run_calc()
        second = deepcopy(first)
        first['rsm'] = first['rsm'].iloc[:0]
        second['name'] = 'Second'
        summary = resume([first, second], 'rsm')
        self.assertEqual(summary.Total.sum(), 35.)

    def test_distributed_template_and_export_without_local_data(self):
        content = (ROOT / 'assets' / 'Template Kriteria Sales Contest.xlsx').read_bytes()
        template = read_template(content)
        self.assertEqual(template['stores']['Store ID'].tolist(), ['CONTOH001', 'CONTOH002'])
        self.assertEqual(template['rsm'].iloc[0]['Reward 100%'], 1000000)
        exported = export_results([self.run_calc()], [content], 'Test')
        workbook = openpyxl.load_workbook(BytesIO(exported), data_only=True)
        self.assertEqual(workbook['Resume AM']['C2'].value, 5.)
        self.assertEqual(workbook['Criteria Kategori bebas']['B4'].value, 'CONTOH001')


if __name__ == '__main__':
    unittest.main()
